# quiz/downloader.py
# Tạo file Word/PDF từ dữ liệu câu hỏi trắc nghiệm.
# Tách từ index.py lines 386–517.

import io
import os
import tempfile
from typing import List

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi import HTTPException


def download_quiz_docx(quiz_data: List[dict], original_filename: str):
    """Xuất đề thi ra file Word (.docx) kèm bảng đáp án."""
    try:
        doc = docx.Document()

        title = doc.add_heading(f'Đề thi từ file: {original_filename}', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for item in quiz_data:
            if isinstance(item.get('question_number'), int):
                q_para = doc.add_paragraph()
                q_para.add_run(f"Câu {item['question_number']}: ").bold = True
                q_para.add_run(item['question_text'])

                for i, option in enumerate(item.get('options', [])):
                    doc.add_paragraph(f"{chr(65 + i)}. {option}")

                doc.add_paragraph()

        # Bảng đáp án
        if any(item.get('correct_answer_char') for item in quiz_data):
            doc.add_heading('Bảng đáp án', level=1)

            table = doc.add_table(rows=1, cols=8)
            table.style = 'Table Grid'

            header_cells = table.rows[0].cells
            for i, header in enumerate(['Câu', 'Đáp án'] * 4):
                header_cells[i].text = header
                header_cells[i].paragraphs[0].runs[0].bold = True

            quiz_with_answers = [q for q in quiz_data if isinstance(q.get('question_number'), int)]
            for i in range(0, len(quiz_with_answers), 4):
                row_cells = table.add_row().cells
                for j in range(4):
                    if i + j < len(quiz_with_answers):
                        item = quiz_with_answers[i + j]
                        row_cells[j * 2].text = str(item['question_number'])
                        row_cells[j * 2 + 1].text = item.get('correct_answer_char', '?')

        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        base_name = original_filename.rsplit('.', 1)[0] if '.' in original_filename else original_filename
        return file_stream, f"{base_name}_quiz.docx"

    except Exception as e:
        print(f"[Downloader] Error creating DOCX: {e}")
        raise HTTPException(status_code=500, detail=f'Lỗi khi tạo file Word: {str(e)}')


def download_quiz_pdf(quiz_data: List[dict], original_filename: str):
    """Xuất đề thi ra file PDF dùng ReportLab."""
    try:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
        except ImportError:
            raise HTTPException(status_code=500, detail='Thư viện reportlab chưa được cài đặt')

        with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp_file:
            pdf_path = tmp_file.name

        c = canvas.Canvas(pdf_path, pagesize=A4)
        width, height = A4
        font_bold = 'Helvetica-Bold'
        font_normal = 'Helvetica'

        def safe_text(text):
            if not text:
                return ""
            try:
                return str(text).encode('utf-8').decode('utf-8')
            except Exception:
                return ''.join(ch for ch in str(text) if ord(ch) < 128)

        # Title
        c.setFont(font_bold, 16)
        title = safe_text(f"Đề thi từ file: {original_filename}")
        title_width = c.stringWidth(title, font_bold, 16)
        c.drawString((width - title_width) / 2, height - 50, title)

        y = height - 100
        c.setFont(font_normal, 12)

        for item in quiz_data:
            if isinstance(item.get('question_number'), int):
                if y < 150:
                    c.showPage()
                    y = height - 50

                q_text = safe_text(f"Câu {item['question_number']}: {item['question_text']}")
                c.drawString(50, y, q_text[:100] + "..." if len(q_text) > 100 else q_text)
                y -= 20

                for i, option in enumerate(item.get('options', [])):
                    if y < 50:
                        c.showPage()
                        y = height - 50
                    opt_text = safe_text(f"{chr(65 + i)}. {option}")
                    c.drawString(70, y, opt_text[:80] + "..." if len(opt_text) > 80 else opt_text)
                    y -= 15

                y -= 10

        c.save()

        with open(pdf_path, 'rb') as f:
            file_stream = io.BytesIO(f.read())
        os.unlink(pdf_path)
        file_stream.seek(0)

        base_name = original_filename.rsplit('.', 1)[0] if '.' in original_filename else original_filename
        return file_stream, f"{base_name}_quiz.pdf"

    except Exception as e:
        print(f"[Downloader] Error creating PDF: {e}")
        raise HTTPException(status_code=500, detail=f'Lỗi khi tạo file PDF: {str(e)}')
