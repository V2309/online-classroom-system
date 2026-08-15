# quiz/extractor.py
# Logic trích xuất câu hỏi từ PDF (PyMuPDF) và DOCX (python-docx).
# Tách từ index.py lines 77–384.

import re
import html
import base64
import fitz  # PyMuPDF
import docx

ALLOWED_EXTENSIONS = {'pdf', 'docx'}


def allowed_file(filename: str) -> bool:
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


# ---------------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------------

def snapshot_drawing(page, drawing_rect, dpi=150) -> str:
    """Chụp ảnh một vùng drawing trong PDF → base64 img HTML."""
    try:
        clip = fitz.Rect(drawing_rect)
        pix = page.get_pixmap(clip=clip, dpi=dpi)
        img_bytes = pix.tobytes("png")
        b64_string = base64.b64encode(img_bytes).decode('utf-8')
        return f'<img src="data:image/png;base64,{b64_string}" style="vertical-align: middle; max-height: 2.5em;" />'
    except Exception as e:
        print(f"Lỗi snapshot: {e}")
        return ""


def group_items_by_lines(items, threshold=5.0) -> list:
    """Gom các items (text/image) theo dòng dựa vào tọa độ Y."""
    lines = []
    if not items:
        return lines
    items.sort(key=lambda item: (item['bbox'].y0, item['bbox'].x0))
    current_line = [items[0]]
    last_y0 = items[0]['bbox'].y0
    for item in items[1:]:
        if abs(item['bbox'].y0 - last_y0) < threshold:
            current_line.append(item)
        else:
            current_line.sort(key=lambda x: x['bbox'].x0)
            lines.append(current_line)
            current_line = [item]
            last_y0 = item['bbox'].y0
    current_line.sort(key=lambda x: x['bbox'].x0)
    lines.append(current_line)
    return lines


def parse_horizontal_options(text: str) -> list:
    """Parse các đáp án A/B/C/D nằm ngang trên 1 dòng."""
    pattern = re.compile(r'([A-D][.:])(.*?)(?=\s*[A-D][.:]|\s*Câu\s*\d+|\Z)', re.IGNORECASE | re.DOTALL)
    matches = pattern.findall(text)
    return [f"{m[0].strip()} {m[1].strip()}" for m in matches]


# ---------------------------------------------------------------------------
# PDF extractor
# ---------------------------------------------------------------------------

def extract_pdf_data(file_path: str) -> list:
    """
    Trích xuất câu hỏi trắc nghiệm từ file PDF.
    Hỗ trợ: text, inline images (drawings), đáp án inline hoặc xuống dòng.
    """
    print(f"\n[Extractor] Bắt đầu xử lý PDF: {file_path}")
    try:
        doc = fitz.open(file_path)
        quiz_data = []
        current_quiz = None
        current_state = "question"
        all_lines = []

        # Thu thập tất cả dòng từ tất cả trang
        for page_num, page in enumerate(doc):
            print(f"[Extractor] --- Trang {page_num + 1}/{len(doc)} ---")
            words = page.get_text("words")
            drawings = page.get_drawings()

            all_items = []
            for w in words:
                bbox = fitz.Rect(w[:4])
                all_items.append({"bbox": bbox, "type": "text", "content": f" {html.escape(w[4])} "})

            for d in drawings:
                bbox = fitz.Rect(d['rect'])
                if bbox.width < 5 or bbox.height < 5:
                    continue
                img_html = snapshot_drawing(page, bbox)
                if img_html:
                    all_items.append({"bbox": bbox, "type": "image", "content": img_html})

            lines_on_page = group_items_by_lines(all_items, threshold=5.0)
            all_lines.extend(lines_on_page)

        print(f"[Extractor] Tổng cộng {len(all_lines)} dòng.")

        option_splitter_regex = re.compile(r'(\s*[A-D][.:])', re.IGNORECASE)

        for line_items in all_lines:
            if not line_items:
                continue

            line_html = "".join([item['content'] for item in line_items]).strip()
            line_text_only = re.sub(r'\s+', ' ', "".join(
                [item['content'] for item in line_items if item['type'] == 'text']
            )).strip()

            if not line_text_only and not line_html:
                continue

            match_question = re.match(r'^(Câu\s*(\d+)[:.]?)', line_text_only, re.IGNORECASE)
            split_by_options_text = option_splitter_regex.split(line_text_only)

            if match_question:
                if current_quiz:
                    quiz_data.append(current_quiz)

                question_number = int(match_question.group(2))
                print(f"    [Extractor] Câu {question_number}")

                question_part_html = line_html
                options_part_html_list = []

                split_by_options_html = option_splitter_regex.split(line_html)
                if len(split_by_options_html) > 1:
                    question_part_html = split_by_options_html[0]
                    for i in range(1, len(split_by_options_html), 2):
                        if i + 1 < len(split_by_options_html):
                            options_part_html_list.append(
                                split_by_options_html[i] + split_by_options_html[i + 1]
                            )

                current_quiz = {
                    "question_number": question_number,
                    "question_text": question_part_html,
                    "options": options_part_html_list,
                    "correct_answer_index": -1,
                    "correct_answer_char": '',
                    "images": [],
                }
                current_state = "option" if options_part_html_list else "question"

            elif current_quiz:
                if len(split_by_options_text) > 1 and not re.match(r'^[A-D][.:]', split_by_options_text[0].strip()):
                    split_by_options_html = option_splitter_regex.split(line_html)
                    first_part = split_by_options_html[0].strip()
                    if first_part:
                        if current_state == "option" and current_quiz["options"]:
                            current_quiz["options"][-1] += " " + first_part
                        else:
                            current_quiz["question_text"] += " " + first_part

                    for i in range(1, len(split_by_options_html), 2):
                        if i + 1 < len(split_by_options_html):
                            current_quiz["options"].append(
                                split_by_options_html[i] + split_by_options_html[i + 1]
                            )
                    current_state = "option"

                elif re.match(r'^[A-D][.:]', line_text_only):
                    current_quiz["options"].append(line_html)
                    current_state = "option"

                else:
                    if current_state == "option" and current_quiz["options"]:
                        current_quiz["options"][-1] += " " + line_html
                    else:
                        current_quiz["question_text"] += " " + line_html

        if current_quiz:
            quiz_data.append(current_quiz)

        doc.close()

        # Dọn dẹp options
        for quiz in quiz_data:
            if quiz["options"]:
                if not quiz["correct_answer_char"]:
                    quiz["correct_answer_char"] = "A"
                    quiz["correct_answer_index"] = 0

                quiz["options"] = [
                    re.sub(r'^(<[^>]+>)*\s*[A-D][.:]\s*', r'\1', opt.strip(), count=1, flags=re.IGNORECASE)
                    for opt in quiz["options"]
                ]

        print(f"[Extractor] Hoàn tất: {len(quiz_data)} câu.")
        return quiz_data

    except Exception as e:
        print(f"[Extractor] Lỗi PDF: {e}")
        import traceback
        traceback.print_exc()
        return []


# ---------------------------------------------------------------------------
# DOCX extractor
# ---------------------------------------------------------------------------

def extract_docx_data(file_path: str) -> list:
    """
    Trích xuất câu hỏi trắc nghiệm từ file DOCX.
    Hỗ trợ: multi-line format và single-line format (fallback).
    """
    try:
        document = docx.Document(file_path)
        quiz_data = []
        current_quiz = None

        for p in document.paragraphs:
            p_text = p.text.strip()
            if not p_text:
                continue

            match_question = re.match(r'^(Câu\s*(\d+)[:.]?)(.*)', p_text, re.IGNORECASE)
            if match_question:
                if current_quiz:
                    quiz_data.append(current_quiz)
                current_quiz = {
                    "question_number": int(match_question.group(2)),
                    "question_text": match_question.group(0),
                    "options": [],
                    "correct_answer_index": -1,
                    "correct_answer_char": '',
                    "images": [],
                }
            elif current_quiz:
                if re.search(r'^[A-D][.:]', p_text):
                    current_quiz["options"].append(p_text.strip())
                else:
                    current_quiz["question_text"] += " " + p_text

        if current_quiz:
            quiz_data.append(current_quiz)

        # Fallback: single-line format
        if not any(q['options'] for q in quiz_data):
            print("[Extractor] DOCX: Fallback to single-line format")
            quiz_data = []
            for p in document.paragraphs:
                p_text = p.text.strip()
                if not p_text or not p_text.startswith("Câu"):
                    continue

                question_blocks = re.split(r'(Câu\s*(\d+)[:.]?)', p_text, flags=re.IGNORECASE)
                if question_blocks[0] == '':
                    question_blocks = question_blocks[1:]

                for i in range(0, len(question_blocks), 3):
                    if i + 2 >= len(question_blocks):
                        continue
                    question_header = question_blocks[i]
                    question_number = int(question_blocks[i + 1])
                    question_body = question_blocks[i + 2].strip()

                    parts = re.split(r'(?=\s*[A-D][.:])', question_body, 1)
                    question_part = parts[0]
                    options_part = parts[1] if len(parts) > 1 else ""
                    options = parse_horizontal_options(options_part)

                    quiz_data.append({
                        "question_number": question_number,
                        "question_text": (question_header + " " + question_part).strip(),
                        "options": options,
                        "correct_answer_index": -1,
                        "correct_answer_char": '',
                        "images": [],
                    })

        # Tìm highlight (đáp án đúng)
        for quiz in quiz_data:
            for p in document.paragraphs:
                if f"Câu {quiz['question_number']}" in p.text:
                    for run in p.runs:
                        if run.font.highlight_color or (run.font.color and run.font.color.rgb):
                            highlighted_text = run.text.strip()
                            for i, option in enumerate(quiz["options"]):
                                if highlighted_text in option:
                                    quiz["correct_answer_char"] = chr(65 + i)
                                    quiz["correct_answer_index"] = i
                                    break

        # Dọn dẹp
        for quiz in quiz_data:
            if quiz["options"]:
                if not quiz["correct_answer_char"]:
                    quiz["correct_answer_char"] = "A"
                    quiz["correct_answer_index"] = 0

                quiz["options"] = [
                    re.sub(r'^[A-D][\.\s]+', '', opt.strip())
                    for opt in quiz["options"]
                ]

            if quiz["correct_answer_char"]:
                quiz["correct_answer_index"] = ord(quiz["correct_answer_char"].upper()) - 65

        return quiz_data

    except Exception as e:
        print(f"[Extractor] Lỗi DOCX: {e}")
        return []
